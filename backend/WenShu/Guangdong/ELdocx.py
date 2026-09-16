"""Fill the official Guangdong evidence list without changing its four-column layout."""

from __future__ import annotations

from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _write_paragraph(paragraph, text: str, *, size: int = 11, align=None):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(str(text or ""))
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    if align is not None:
        paragraph.alignment = align


def _write_cell(cell, text: str, *, align=WD_ALIGN_PARAGRAPH.CENTER, size: int = 10):
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        cell._tc.remove(extra._p)
    _write_paragraph(paragraph, text, size=size, align=align)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _format_date(value: str) -> str:
    try:
        parsed = date.fromisoformat(value)
        return f"{parsed.year}年{parsed.month}月{parsed.day}日"
    except (TypeError, ValueError):
        return str(value or "")


def generate_ELdocx(data: dict, template_docx: str | Path):
    doc = Document(str(template_docx))
    table = doc.tables[0]
    items = [
        item for item in data.get("evidenceList", [])
        if any(str(item.get(key, "")).strip() for key in ("name", "pageCount", "source", "purpose"))
    ]
    if not items and not any(
        str(data.get(key, "")).strip()
        for key in ("committeeName", "caseNumber", "submitterName", "submitDate")
    ):
        return doc

    if len(items) > 9:
        template_row = table.rows[-1]._tr
        for _ in range(len(items) - 9):
            table._tbl.append(deepcopy(template_row))

    for index, item in enumerate(items):
        row = table.rows[index + 1]
        proof = f"来源：{item.get('source', '')}；{item.get('purpose', '')}".strip("；")
        _write_cell(row.cells[0], str(index + 1))
        _write_cell(row.cells[1], item.get("name", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
        _write_cell(row.cells[2], proof, align=WD_ALIGN_PARAGRAPH.LEFT)
        _write_cell(row.cells[3], item.get("pageCount", ""))

    if data.get("submitterName") or data.get("submitDate"):
        footer = f"提交人：{data.get('submitterName', '')}        提交时间：{_format_date(data.get('submitDate', ''))}"
        _write_paragraph(doc.paragraphs[3], footer, size=12)
    return doc
