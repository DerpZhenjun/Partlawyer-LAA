"""Fill the official Zhejiang worker arbitration application template."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _write_cell(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, size: int = 11):
    """Replace a cell's text while preserving cell geometry and paragraph properties."""
    paragraphs = cell.paragraphs
    paragraph = paragraphs[0]
    for extra in paragraphs[1:]:
        cell._tc.remove(extra._p)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    run = paragraph.add_run(str(text or ""))
    run.font.name = "仿宋_GB2312"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    run.font.size = Pt(size)
    run.bold = bold
    paragraph.alignment = align
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _format_date(value: str) -> str:
    try:
        parsed = date.fromisoformat(value)
        return f"{parsed.year}年{parsed.month}月{parsed.day}日"
    except (TypeError, ValueError):
        return str(value or "")


def _household_options(selected: str) -> str:
    options = ["本省非农业户口", "本省农业户口", "外省非农业户口", "外省农业户口", "港澳台人员", "外籍人员"]
    return "  ".join(f"{'√' if item == selected else ' '} {item}" for item in options)


def _keep_worker_form_only(doc):
    """The official file contains worker and employer forms; keep the worker form only."""
    body = doc._element.body
    first_table_seen = False
    for child in list(body):
        if child.tag == qn("w:tbl") and not first_table_seen:
            first_table_seen = True
            continue
        if first_table_seen and child.tag != qn("w:sectPr"):
            body.remove(child)


def generate_LAdocx(data: dict, template_docx: str | Path):
    doc = Document(str(template_docx))
    _keep_worker_form_only(doc)
    table = doc.tables[0]

    requests = [str(item).strip() for item in data.get("requests", []) if str(item).strip()]
    content_fields = (
        "committeeName", "applicantName", "birthDate", "gender", "nationality", "idNumber",
        "phone", "householdType", "idAddress", "respondentName", "socialCreditCode",
        "registeredAddress", "mailingAddress", "workAddress", "legalRepName", "legalRepPhone",
        "legalRepPosition", "companyContact", "companyContactPhone", "factsReason", "applicationDate",
    )
    if not requests and not any(str(data.get(key, "")).strip() for key in content_fields):
        return doc

    _write_cell(table.rows[1].cells[1], data.get("applicantName", ""))
    _write_cell(table.rows[1].cells[4], _format_date(data.get("birthDate", "")), align=WD_ALIGN_PARAGRAPH.CENTER)
    _write_cell(table.rows[2].cells[1], data.get("gender", ""), align=WD_ALIGN_PARAGRAPH.CENTER)
    _write_cell(table.rows[2].cells[4], data.get("nationality", ""), align=WD_ALIGN_PARAGRAPH.CENTER)
    _write_cell(table.rows[3].cells[1], data.get("idNumber", ""))
    _write_cell(table.rows[3].cells[4], data.get("phone", ""), align=WD_ALIGN_PARAGRAPH.CENTER)
    if data.get("householdType"):
        _write_cell(table.rows[4].cells[1], _household_options(data.get("householdType", "")), size=9)
    _write_cell(table.rows[5].cells[2], data.get("idAddress", ""))

    _write_cell(table.rows[7].cells[2], data.get("respondentName", ""))
    _write_cell(table.rows[8].cells[2], data.get("socialCreditCode", ""))
    _write_cell(table.rows[9].cells[2], data.get("registeredAddress", ""))
    _write_cell(table.rows[10].cells[2], data.get("mailingAddress", ""))
    _write_cell(table.rows[11].cells[2], data.get("workAddress", ""))
    _write_cell(table.rows[12].cells[2], data.get("legalRepName", ""))
    _write_cell(table.rows[12].cells[4], data.get("legalRepPhone", ""), align=WD_ALIGN_PARAGRAPH.CENTER)
    _write_cell(table.rows[13].cells[2], data.get("legalRepName", ""))
    _write_cell(table.rows[13].cells[4], data.get("legalRepPosition", ""), align=WD_ALIGN_PARAGRAPH.CENTER)
    _write_cell(table.rows[14].cells[2], data.get("companyContact", ""))
    _write_cell(table.rows[14].cells[4], data.get("companyContactPhone", ""), align=WD_ALIGN_PARAGRAPH.CENTER)

    for offset, row_index in enumerate((16, 17, 18, 19)):
        if offset < len(requests):
            _write_cell(table.rows[row_index].cells[0], f"{offset + 1}．{requests[offset]}", size=11)

    _write_cell(table.rows[21].cells[0], data.get("factsReason", ""), size=11)

    apply_date = _format_date(data.get("applicationDate", "")) or "    年   月   日"
    footer = (
        "此致\n敬礼\n"
        f"{data.get('committeeName', '')}\n"
        f"申请人：{data.get('applicantName', '')}\n"
        f"{apply_date}\n"
        f"附：申请书副本 {data.get('copyCount', 1)} 份。\n"
        "注：1.申请书应用黑色钢笔、中性笔书写或打印。\n"
        "2.申请人应同时提交身份证复印件或其他身份证件。\n"
        "3.事实与理由部分空格不够用时，可用A4纸续加中页。\n"
        "4.申请书副本份数，应按对方当事人人数提交。"
    )
    _write_cell(table.rows[22].cells[0], footer, size=10)
    return doc
