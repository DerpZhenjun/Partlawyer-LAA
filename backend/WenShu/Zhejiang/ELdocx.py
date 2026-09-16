"""Fill the official Zhejiang evidence-list template."""

from __future__ import annotations

from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _write_cell(cell, text: str, *, align=WD_ALIGN_PARAGRAPH.CENTER, size: int = 10):
    paragraphs = cell.paragraphs
    paragraph = paragraphs[0]
    for extra in paragraphs[1:]:
        cell._tc.remove(extra._p)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(str(text or ""))
    run.font.name = "仿宋_GB2312"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    run.font.size = Pt(size)
    paragraph.alignment = align
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _format_date(value: str) -> str:
    try:
        parsed = date.fromisoformat(value)
        return f"{parsed.year}年{parsed.month}月{parsed.day}日"
    except (TypeError, ValueError):
        return str(value or "")


def generate_ELdocx(data: dict, template_docx: str | Path):
    doc = Document(str(template_docx))
    table = doc.tables[0]

    if data.get("committeeName"):
        paragraph = doc.paragraphs[0]
        paragraph.text = data["committeeName"]

    _write_cell(table.rows[0].cells[1], data.get("caseNumber", ""))
    _write_cell(table.rows[0].cells[5], str(data.get("copyCount", 2)))

    items = [
        item for item in data.get("evidenceList", [])
        if any(str(item.get(key, "")).strip() for key in ("name", "pageCount", "source", "purpose"))
    ]
    if not items and not any(
        str(data.get(key, "")).strip()
        for key in ("committeeName", "caseNumber", "submitterName", "submitDate")
    ):
        return doc

    blank_rows = 9
    if len(items) > blank_rows:
        signature_row = table.rows[11]._tr
        template_row = table.rows[10]._tr
        for _ in range(len(items) - blank_rows):
            signature_row.addprevious(deepcopy(template_row))

    for index in range(max(blank_rows, len(items))):
        row = table.rows[2 + index]
        item = items[index] if index < len(items) else {}
        _write_cell(row.cells[0], str(index + 1) if item else "")
        _write_cell(row.cells[1], item.get("name", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
        _write_cell(row.cells[3], item.get("pageCount", ""))
        _write_cell(row.cells[4], item.get("source", ""))
        _write_cell(row.cells[5], item.get("purpose", ""), align=WD_ALIGN_PARAGRAPH.LEFT)

    submit_row = table.rows[2 + max(blank_rows, len(items))]
    _write_cell(submit_row.cells[2], data.get("submitterName", ""))
    _write_cell(submit_row.cells[5], _format_date(data.get("submitDate", "")))

    # The source template's 14 pt footer can leave the final two characters on a
    # second page in Word. Keep the official wording but size it to the A4 page.
    footer = doc.paragraphs[-1]
    footer.text = f"本表一式{data.get('copyCount', 2)}份提交仲裁委。证据原件请提交人妥善保管，在开庭时提交核对。"
    footer.paragraph_format.space_before = Pt(3)
    footer.paragraph_format.space_after = Pt(0)
    for run in footer.runs:
        run.font.name = "仿宋_GB2312"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
        run.font.size = Pt(11)
    return doc
